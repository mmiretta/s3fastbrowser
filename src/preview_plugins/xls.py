"""Excel workbooks (.xlsx, .xlsm) as tables, one tab per sheet (first 200
rows and 50 columns of each). Legacy binary .xls files are not supported by
openpyxl and fall through to the next plugin."""
from io import BytesIO

from openpyxl import load_workbook

from preview_plugins._util import esc, note

NAME = 'xlsx'
MAX_BYTES = None
MAX_ROWS = 200
MAX_COLS = 50
MAX_SHEETS = 10


def can_preview(filename, content_type):
    return filename.lower().endswith(('.xlsx', '.xlsm')) or content_type == \
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def render_preview(fileobj, filename, content_type, bucket_name=None):
    try:
        workbook = load_workbook(BytesIO(fileobj.read()), read_only=True, data_only=True)
    except Exception as e:
        return note(f'Could not open the workbook: {esc(e)}', 'warn')

    tabs, panes = [], []
    for idx, name in enumerate(workbook.sheetnames[:MAX_SHEETS]):
        sheet = workbook[name]
        rows = []
        for r, row in enumerate(sheet.iter_rows(values_only=True)):
            if r >= MAX_ROWS:
                break
            rows.append(['' if v is None else str(v) for v in row[:MAX_COLS]])
        html = ['<div class="pv-table-wrap"><table class="table table-sm table-bordered pv-table">']
        if rows:
            html.append('<thead><tr>' + ''.join(f'<th>{esc(c)}</th>' for c in rows[0]) + '</tr></thead><tbody>')
            for row in rows[1:]:
                html.append('<tr>' + ''.join(f'<td>{esc(c)}</td>' for c in row) + '</tr>')
            html.append('</tbody>')
        else:
            html.append('<tbody><tr><td>(empty sheet)</td></tr></tbody>')
        html.append('</table></div>')
        active = ' active' if idx == 0 else ''
        tabs.append(f'<li class="nav-item"><button class="nav-link{active}" data-bs-toggle="tab" '
                    f'data-bs-target="#pv-sheet-{idx}" type="button">{esc(name)}</button></li>')
        panes.append(f'<div class="tab-pane fade{" show active" if idx == 0 else ""}" id="pv-sheet-{idx}">'
                     + ''.join(html) + '</div>')
    workbook.close()
    return (f'<ul class="nav nav-tabs mb-2">{"".join(tabs)}</ul><div class="tab-content">{"".join(panes)}</div>'
            + note(f'Showing up to {MAX_ROWS} rows and {MAX_COLS} columns per sheet.'))
